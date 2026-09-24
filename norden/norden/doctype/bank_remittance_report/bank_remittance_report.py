# Copyright (c) 2024, Teampro and contributors
# For license information, please see license.txt
import frappe
from frappe.model.document import Document
from datetime import date, timedelta, datetime,time
from frappe.utils import (getdate, cint, add_months, date_diff, add_days,
    nowdate, get_datetime_str, cstr, get_datetime, now_datetime, format_datetime,today, format_date)
import math
from datetime import datetime
import openpyxl
import calendar
from openpyxl.styles import Alignment, Font, PatternFill
from io import BytesIO
import frappe
from frappe.model.document import Document


class BankRemittanceReport(Document):
	pass

@frappe.whitelist()
def download():
    frappe.errprint("Hello")
    filename = 'Bank Remittance Report'
    vdate = None  
    company = None  
    test = build_xlsx_response(filename, vdate, company)


def get_emp_data(vdate,company):
    employee = frappe.get_list('Employee',{'status':'Active','company':company},['*'])
    filter_values = []
    for emp in employee:
        filter_values.append(emp)
    return filter_values

# def date_formatting(vdate):
     

def make_xlsx(filename, vdate, company):
    args = frappe.local.form_dict
    emp_details = get_emp_data(args.vdate, args.company)
    value_date=args.vdate
    year, month, date = value_date.split('-')
    month_name = calendar.month_name[int(month)]
    current_year = int(year)
    current_month = month_name
    salary_statement = f'Salary {month_name} {current_year}'
    frappe.errprint(current_year)
    frappe.errprint(current_month)
    acctype=''
    date_obj = datetime.strptime(value_date, '%Y-%m-%d')
    formatted_date = date_obj.strftime('%d/%m/%Y')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = filename 
    headers = [
    "Transaction Type", "Beneficiary Code","Beneficiary Acc No","Amount", "Beneficiary Name", 
    "Drawee Location", "Print Location", "Bene Add 1", "Bene Add 2", 
    "Bene Add 3", "Bene Add 4", "Bene Add 5", "Instruction Reference", 
    "Customer Reference",
    "Pay Details 1", "Pay Details 2", "Pay Details 3", "Pay Details 4", 
    "Pay Details 5", "Pay Details 6", "Pay Details 7", "Cheque Number", 
    "Value Date", "MICR Number", "IFCS Code", "Bene Bank", "Bene Branch", 
    "Email ID"
	]
    ws.append(headers)
    align_center = Alignment(horizontal='center', vertical='center')
    header_row = ws[1]  
    header_font = Font(color="FFFFFF")  
    header_fill = PatternFill(start_color='CF0243', end_color='CF0243', fill_type='solid')  # Red color
    for cell in header_row:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
    for i in emp_details:
        if i['bank_name']=='HDFC Bank':
            acctype = 'I'
        else:
            acctype = 'N'
        ws.append([
            acctype,
            i['custom_beneficiary_name'], #Beneficiary code
            i['bank_ac_no'],'',i['custom_beneficiary_name'],'',#amount and bene name
            '','','','','','','',
            salary_statement,'',#string value
            '','','','','','','',formatted_date,
            '',i['ifsc_code'],i['bank_name'],i['bank_branch'],
            i['user_id']
        ])
    for col in ws.iter_cols(min_col=1, max_col=len(headers)):
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2) * 1.2
        ws.column_dimensions[column].width = adjusted_width
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    return xlsx_file

def build_xlsx_response(filename, vdate, company):
    xlsx_file = make_xlsx(filename, vdate, company)
    frappe.local.response.filename = filename + '.xlsx'
    frappe.local.response.filecontent = xlsx_file.getvalue()
    frappe.local.response.type = "binary"


