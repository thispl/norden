# Copyright (c) 2022, Abdulla and contributors
# For license information, please see license.txt

import frappe
from frappe.model.document import Document
from frappe.utils.xlsxutils import read_xlsx_file_from_attached_file
from frappe.utils.file_manager import get_file

from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
import frappe
from frappe.utils import flt
from jinja2 import Template

import openpyxl
import xlrd
import re
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO, string_types
from frappe.utils import (
    flt,
    cint,
    cstr,
    get_html_format,
    get_url_to_form,
    gzip_decompress,
    format_duration,
)

class AllocationDetails(Document):
	@frappe.whitelist()
	def get_data(self):
		date = frappe.db.get_value("Custom Settings","Custom Settings","date")
		data = ''
		data1 = ''
		item = frappe.get_value('Item',{'item_code':self.item_code},'item_code')
		rate = frappe.get_value('Item',{'item_code':self.item_code},'valuation_rate')
		group = frappe.get_value('Item',{'item_code':self.item_code},'item_group')
		des = frappe.get_value('Item',{'item_code':self.item_code},'description')
		price = frappe.get_value('Item Price',{'item_code':self.item_code,'price_list':'Cost'},'price_list_rate')
		c_s_p = frappe.get_value('Item Price',{'item_code':self.item_code,'price_list':'Standard Selling'},'price_list_rate') or 0
		csp = 'Current Selling Price'
		cpp = 'Current Purchase Price'
		cost = 'COST'
		pso = 'Pending Sales order'
		po ='Total Purchase Order'
		ppo = 'Pending Purchase order'
		cspp_rate = 0
		cppp_rate = 0
		psoc = 0
		ppoc = 0
		ppoc_total = 0
		i = 0
		cou = 0
		p_po = 0
		p_so = 0
		tot = 'Total'
		uom = 'Nos'

		# stocks_query = frappe.db.sql("""select actual_qty,reserved_stock,warehouse,stock_uom,stock_value from tabBin
		# 		where item_code = '%s' """%(item),as_dict=True)
		stocks_query = frappe.db.sql("""
			select b.actual_qty, b.reserved_stock, b.warehouse, b.stock_uom, b.stock_value
			from tabBin b
			join tabWarehouse w on w.name = b.warehouse
			where b.item_code = %s
			and w.custom_is_service_stock = 0
		""", (item,), as_dict=True)
		# frappe.errprint(stocks_query.reserved_stock)
		if stocks_query:
			stocks = stocks_query
		
		data += '<table class="table table-bordered" style="width:70%"><tr><th style="padding:1px;border: 1px solid black;background-color:#fe3f0c;color:white" colspan=10><center>NORDEN PRODUCT SEARCH</center></th></tr>'
		data += '<tr><td colspan = 2 style="padding:1px;border: 1px solid black;color:white;background-color:#6f6f6f;text-align: left"><b>Item Code</b></td><td colspan = 8 style="padding:1px;border: 1px solid black;text-align: left"><b>%s</b></td></tr>'%(item)
		data += '<tr><td colspan = 2 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;text-align: left"><b>Item Name</b></td><td colspan = 8 style="padding:1px;border: 1px solid black;text-align: left"><b>%s</b></td></tr>'%(frappe.db.get_value('Item',item,'item_name'))
		data += '<tr><td colspan = 2 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;text-align: left"><b>Item Group</b></td><td colspan = 8 style="padding:1px;border: 1px solid black;text-align: left"><b>%s</b></td></tr>'%(group)
		data += '<tr><td colspan = 2 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;text-align: left"><b>Description</b></td><td colspan = 8 style="padding:1px;border: 1px solid black;text-align: left"><b>%s</b></td></tr>'%(des)
		
		if stocks_query:
			data += '<tr><td colspan = 1 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;"><center><b>Company</b></center></td>'
			data += '<td colspan = 1 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;"><center><b>Warehouse</b></center></td>'
			data += '<td colspan = 1 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;"><center><b>QTY</b></center></td>'
			data += '<td colspan = 1 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;"><center><b>UOM</b></center></td>'
			data += '<td colspan = 1 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;"><center><b>Pending PO</b></center></td>'
			data += '<td colspan = 1 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;"><center><b>Pending SO</b></center></td>'
			data += '</tr>'
			for stock in stocks_query:
				if stock.actual_qty >= 0:
					reserve = stock.actual_qty - stock.reserved_stock
					# stock_company = frappe.db.sql("""select company from tabWarehouse where name = '%s' and custom_stock_is_shown_only_in_product_search = 0 and company = "Norden Communication Middle East FZE" """%(stock.warehouse),as_dict=True)
					stock_company = frappe.db.sql("""
						select company from tabWarehouse 
						where name = %s 
						and custom_stock_is_shown_only_in_product_search = 0
						and custom_is_service_stock = 0
						and company = "Norden Communication Middle East FZE"
					""", (stock.warehouse,), as_dict=True)
					for com in stock_company:
						if com.company != "Norden Africa":
							new_so = frappe.db.sql("""select sum(`tabSales Order Item`.qty) as qty,sum(`tabSales Order Item`.delivered_qty) as d_qty from `tabSales Order`
							left join `tabSales Order Item` on `tabSales Order`.name = `tabSales Order Item`.parent
							where `tabSales Order Item`.item_code = '%s' and `tabSales Order`.docstatus = 1 and `tabSales Order`.status != 'Closed' and `tabSales Order`.company = '%s' and `tabSales Order`.set_warehouse = '%s' and `tabSales Order`.transaction_date >= '%s'  """ % (self.item_code,com.company,stock.warehouse,date), as_dict=True)[0]
							if not new_so['qty']:
								new_so['qty'] = 0
							if not new_so['d_qty']:
								new_so['d_qty'] = 0
							del_total = new_so['qty'] - new_so['d_qty']

							new_po = frappe.db.sql("""select sum(`tabPurchase Order Item`.qty) as qty,sum(`tabPurchase Order Item`.received_qty) as d_qty from `tabPurchase Order`
							left join `tabPurchase Order Item` on `tabPurchase Order`.name = `tabPurchase Order Item`.parent
							where `tabPurchase Order Item`.item_code = '%s' and `tabPurchase Order`.docstatus = 1 and `tabPurchase Order`.status != 'Closed' and `tabPurchase Order`.company = '%s' and `tabPurchase Order`.set_warehouse = '%s' and `tabPurchase Order`.transaction_date >= '%s'  """ % (self.item_code,com.company,stock.warehouse,date), as_dict=True)[0]
							if not new_po['qty']:
								new_po['qty'] = 0
							if not new_po['d_qty']:
								new_po['d_qty'] = 0
							ppoc_total = new_po['qty'] - new_po['d_qty']

							country,default_currency = frappe.get_value("Company",{"name":com.company},["country","default_currency"])
							if country == "United Arab Emirates":
								cost = frappe.get_value("Item Price",{"item_code":item,"price_list":"Electra Qatar - NCMEF"},["price_list_rate"])
							else:
								cost = frappe.get_value("Item Price",{"item_code":item,"price_list":"STANDARD BUYING-USD"},["price_list_rate"])
							
							valuation_rate = 0
							source_warehouse = frappe.db.get_value('Warehouse', {'default_for_stock_transfer':1,'company': com.company }, ["name"])
							latest_vr = frappe.db.sql("""select valuation_rate as vr from tabBin
									where item_code = '%s' and warehouse = '%s' """%(item,source_warehouse),as_dict=True)
							if latest_vr:
								valuation_rate = latest_vr[0]["vr"]
							else:
								val_rate = []
								l_vr = frappe.db.sql("""
									SELECT valuation_rate AS vr FROM tabBin
									WHERE item_code = %s AND warehouse = %s
								""", (item, source_warehouse), as_dict=True)
								for item in l_vr: 
									if item not in val_rate: 
										val_rate.append(item.vr)
								if len(val_rate) > 1 :
									valuation_rate = max(val_rate)
							
							
							pricelist = country + ' ' + "Sales Price"
							if country == "United Arab Emirates":
								sp = frappe.get_value("Item Price",{"item_code":item,"price_list":"Internal - NCMEF"},["price_list_rate"])
							else:
								sp = frappe.get_value("Item Price",{"item_code":item,"price_list":pricelist},["price_list_rate"])
							data += '<tr><td colspan = 1 style="padding:1px;border: 1px solid black">%s</td><td colspan = 1 style="padding:1px;border: 1px solid black">%s</td><td colspan = 1 style="padding:1px;border: 1px solid black"><center><b>%s</b></center></td><td colspan = 1 style="padding:1px;border: 1px solid black"><center><b>%s</b></center></td><td colspan = 1 style="padding:1px;border: 1px solid black"><center><b>%s</b></center></td><td colspan = 1 style="padding:1px;border: 1px solid black"><center><b>%s</b></center></td></tr>'%(com.company,stock.warehouse,int(reserve) or 0,stock.stock_uom or '-',int(ppoc_total) or 0,int(del_total) or 0)
							i += 1
							# cou += stock.actual_qty
							cou += reserve
							p_po += ppoc_total
							p_so += del_total
			data += '<tr><td align="right" colspan = 2 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;"><b>%s</b></td><td colspan = 1 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;"><center><b>%s</b></center></td><td colspan = 1 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;"><center><b>%s</b></center></td><td colspan = 1 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;"><center><b>%s</b></center></td><td colspan = 1 style="padding:1px;border: 1px solid black;background-color:#6f6f6f;color:white;"><center><b>%s</b></center></td></tr>'%(tot or 0,int(cou) or 0,uom,p_po or 0,p_so or 0)
			data += '</table>'
		else:
			i += 1
			data1 += '<tr><td align="center" colspan = 10 style="padding:1px;border: 1px solid black;background-color:#fe3f0c;color:white;"><b>No Stock Available</b></td></tr>'
			data1 += '</table>'
			data += data1
		if i > 0:
			return data

	@frappe.whitelist()
	def get_item_summary(self):
		date = frappe.db.get_value("Custom Settings", "Custom Settings", "date")
		data = []

		if self.like:
			like_filter = "%" + self.like + "%"
			item = frappe.get_all("Item", {"name": ["like", like_filter]}, ["*"])
		elif self.item_code:
			item = frappe.get_all("Item", {"name": self.item_code}, ["*"])
		
		else:
			return "<p>No item selected.</p>"

		for i in item:
			# 1. Actual Stock
			# stocks = frappe.db.sql("""
			# 	SELECT (SUM(`tabBin`.actual_qty) - SUM(reserved_stock)) AS actual_qty
			# 	FROM `tabBin`
			# 	JOIN `tabWarehouse` ON `tabWarehouse`.name = `tabBin`.warehouse
			# 	WHERE `tabBin`.item_code = %s AND `tabWarehouse`.company = %s AND disabled = 0 and `tabWarehouse`.custom_stock_is_shown_only_in_product_search = 0
			# """, (i.name, self.company), as_dict=True)[0]
			stocks = frappe.db.sql("""
				SELECT (SUM(`tabBin`.actual_qty) - SUM(reserved_stock)) AS actual_qty
				FROM `tabBin`
				JOIN `tabWarehouse` ON `tabWarehouse`.name = `tabBin`.warehouse
				WHERE `tabBin`.item_code = %s
				AND `tabWarehouse`.company = %s
				AND disabled = 0
				AND `tabWarehouse`.custom_stock_is_shown_only_in_product_search = 0
				AND `tabWarehouse`.custom_is_service_stock = 0
			""", (i.name, self.company), as_dict=True)[0]

			stocks["actual_qty"] = stocks["actual_qty"] or 0

			# 2. Reserved Quantity
			total_reserved_qty = 0
			stock_entries = frappe.get_all(
				"Stock Reservation Entry",
				filters={
					"company": self.company,
					"item_code": i.name,
					"status": ["in", ["Reserved", "Partially Reserved"]]
				},
				fields=["reserved_qty","voucher_no"]
			)
			for entry in stock_entries:

				total_reserved_qty += entry.reserved_qty or 0

			prq = frappe.db.sql("""
				SELECT (reserved_qty - delivered_qty) AS partially_delivered
				FROM `tabStock Reservation Entry`
				WHERE item_code = %s AND company = %s AND status = 'Partially Delivered'
			""", (i.name, self.company), as_dict=True)
			for prq_ in prq:
				total_reserved_qty += flt(prq_['partially_delivered']) or 0

			# 3. Non-sale (MRB) Quantity
			mrb_qty = frappe.db.sql("""
				SELECT (SUM(`tabBin`.actual_qty) - SUM(reserved_stock)) AS actual_qty
				FROM `tabBin`
				JOIN `tabWarehouse` ON `tabWarehouse`.name = `tabBin`.warehouse
				WHERE `tabBin`.item_code = %s AND `tabWarehouse`.company = %s AND `tabWarehouse`.custom_is_non_sale = 1
			""", (i.name, self.company), as_dict=True)[0]
			mrb_qty["actual_qty"] = mrb_qty["actual_qty"] or 0

			# 4. Demo Qty
			demo_qty = frappe.db.sql("""
				SELECT (SUM(`tabBin`.actual_qty) - SUM(reserved_stock)) AS actual_qty
				FROM `tabBin`
				JOIN `tabWarehouse` ON `tabWarehouse`.name = `tabBin`.warehouse
				WHERE `tabBin`.item_code = %s AND `tabWarehouse`.company = %s AND custom_is_demo = 1
			""", (i.name, self.company), as_dict=True)[0]
			demo_qty["actual_qty"] = demo_qty["actual_qty"] or 0

			# 5. PO Qty
			new_po = frappe.db.sql("""
				SELECT SUM(`tabPurchase Order Item`.qty) AS qty, SUM(`tabPurchase Order Item`.received_qty) AS d_qty
				FROM `tabPurchase Order`
				LEFT JOIN `tabPurchase Order Item` ON `tabPurchase Order`.name = `tabPurchase Order Item`.parent
				WHERE `tabPurchase Order Item`.item_code = %s AND `tabPurchase Order`.docstatus = 1
				AND `tabPurchase Order`.company = %s AND `tabPurchase Order Item`.qty >= `tabPurchase Order Item`.received_qty
				AND `tabPurchase Order`.transaction_date >= %s
			""", (i.name, self.company, date), as_dict=True)[0]
			ppoc_total = (new_po['qty'] or 0) - (new_po['d_qty'] or 0)

			# 6. SO Pending Qty
			sa = frappe.db.sql("""
				SELECT `tabSales Order Item`.parent AS parent, SUM(`tabSales Order Item`.qty) AS qty,
					SUM(`tabSales Order Item`.delivered_qty) AS delivered_qty
				FROM `tabSales Order`
				LEFT JOIN `tabSales Order Item` ON `tabSales Order`.name = `tabSales Order Item`.parent
				WHERE `tabSales Order Item`.item_code = %s AND `tabSales Order`.docstatus != 2
				AND `tabSales Order`.company = %s AND `tabSales Order`.transaction_date >= %s AND `tabSales Order`.status != 'Closed'
				AND `tabSales Order`.per_delivered != 100
				GROUP BY `tabSales Order Item`.parent
			""", (i.name, self.company, date), as_dict=True)

			psoc_total = 0
			for j in sa:
				psoc_total += (j.qty or 0) - (j.delivered_qty or 0)

			# 7. Scrap warehouse adjustment
			total_scrap_qty = 0
			if stocks["actual_qty"] > 0:
				ware = frappe.db.sql("""
					SELECT name FROM `tabWarehouse`
					WHERE is_scrap = 1 AND disabled = 0 AND company = %s
				""", (self.company,), as_dict=True)
				for house in ware:
					sc_qty = frappe.get_value("Bin", {"warehouse": house.name, "item_code": i.name}, "actual_qty") or 0
					total_scrap_qty += sc_qty

			if stocks["actual_qty"] > 0:
				row = {
					'company': self.company,
					'item': i.name,
					'item_name': i.item_name,
					'total_qty': stocks["actual_qty"] - total_scrap_qty + total_reserved_qty,
					'reserved_qty': total_reserved_qty,
					'free_qty': stocks["actual_qty"] - total_scrap_qty - demo_qty['actual_qty']-mrb_qty['actual_qty'],
					'non_sale_qty': mrb_qty['actual_qty'],
					'po_qty': ppoc_total,
					'so_qty': psoc_total,
					'unit': i.stock_uom,
					'demo_qty': demo_qty['actual_qty']
				}
			else:
				row = {
					'company': self.company,
					'item': i.name,
					'item_name': i.item_name,
					'total_qty': 0 + total_reserved_qty,
					'reserved_qty': total_reserved_qty,
					'free_qty': 0,
					'non_sale_qty': mrb_qty['actual_qty'],
					'po_qty': ppoc_total,
					'so_qty': psoc_total,
					'unit': i.stock_uom,
					'demo_qty': demo_qty['actual_qty']
				}
			data.append(row)

		return data